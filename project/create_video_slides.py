#!/usr/bin/env python3
"""Create the nine-scene visual deck for the narrated project video."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SCRIPT_PATH = DOCS / "VIDEO_SCRIPT_8MIN.md"
OUTPUT_PATH = DOCS / "Workload_Aware_Chiplet_Video_Deck.pptx"

BG = RGBColor(7, 17, 31)
PANEL = RGBColor(15, 32, 52)
PANEL_2 = RGBColor(22, 43, 67)
WHITE = RGBColor(246, 249, 252)
MUTED = RGBColor(171, 187, 202)
GRID = RGBColor(83, 105, 127)
MESH = RGBColor(94, 168, 255)
FIXED = RGBColor(255, 183, 77)
AWARE = RGBColor(48, 213, 200)
CORAL = RGBColor(255, 107, 107)
PURPLE = RGBColor(178, 132, 255)
GREEN = RGBColor(105, 219, 124)

SW = Inches(13.333)
SH = Inches(7.5)


def rect(slide, x, y, w, h, fill, radius=True, line=None, line_width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def textbox(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=WHITE,
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Aptos",
    margin=0.03,
    valign=MSO_ANCHOR.MIDDLE,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=GRID, width=2):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def pill(slide, text, x, y, w, fill, color=BG, size=11):
    rect(slide, x, y, w, 0.32, fill)
    textbox(slide, text, x + 0.04, y, w - 0.08, 0.32, size, color, True, PP_ALIGN.CENTER)


def base_slide(prs, number, timing, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = BG
    pill(slide, f"SCENE {number}", 0.52, 0.28, 0.86, AWARE)
    pill(slide, timing, 1.48, 0.28, 1.10, PANEL_2, WHITE)
    textbox(slide, title, 0.52, 0.72, 12.25, 0.62, 26, WHITE, True)
    if subtitle:
        textbox(slide, subtitle, 0.54, 1.27, 12.0, 0.36, 12, MUTED)
    textbox(
        slide,
        f"Scene {number}  •  Narration: docs/VIDEO_SCRIPT_8MIN.md",
        0.54,
        7.12,
        12.2,
        0.22,
        9,
        MUTED,
    )
    return slide


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def draw_mesh(slide, x, y, size, highlight=None, shortcut=None, labels=True):
    gap = size / 3
    positions = {r * 4 + c: (x + c * gap, y + r * gap) for r in range(4) for c in range(4)}
    highlight_edges = set()
    if highlight:
        highlight_edges = {
            tuple(sorted((highlight[index], highlight[index + 1])))
            for index in range(len(highlight) - 1)
        }
    for r in range(4):
        for c in range(4):
            node = r * 4 + c
            if c < 3:
                other = node + 1
                edge = tuple(sorted((node, other)))
                line(slide, *positions[node], *positions[other], CORAL if edge in highlight_edges else GRID, 4 if edge in highlight_edges else 2)
            if r < 3:
                other = node + 4
                edge = tuple(sorted((node, other)))
                line(slide, *positions[node], *positions[other], CORAL if edge in highlight_edges else GRID, 4 if edge in highlight_edges else 2)
    if shortcut:
        line(slide, *positions[shortcut[0]], *positions[shortcut[1]], AWARE, 5)
    node_size = 0.33 if size > 2.3 else 0.20
    for node, (nx, ny) in positions.items():
        fill = WHITE
        if highlight and node in highlight:
            fill = CORAL
        if shortcut and node in shortcut:
            fill = AWARE
        rect(slide, nx - node_size / 2, ny - node_size / 2, node_size, node_size, fill, True)
        if labels:
            textbox(
                slide,
                str(node),
                nx - node_size / 2,
                ny - node_size / 2,
                node_size,
                node_size,
                9 if size > 2.3 else 6,
                BG,
                True,
                PP_ALIGN.CENTER,
                margin=0,
            )
    return positions


def add_picture_contain(slide, path, x, y, w, h, pad=0.04):
    path = Path(path)
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    pw, ph = iw * scale, ih * scale
    left = x + (w - pw) / 2
    top = y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(pw), Inches(ph))


def label_card(slide, title, value, x, y, w, color):
    rect(slide, x, y, w, 0.90, PANEL, True, color, 2)
    textbox(slide, title, x + 0.12, y + 0.08, w - 0.24, 0.25, 10, MUTED, True)
    textbox(slide, value, x + 0.12, y + 0.34, w - 0.24, 0.42, 19, color, True)


def extract_narration() -> dict[str, str]:
    text = SCRIPT_PATH.read_text(encoding="utf-8")
    pattern = re.compile(
        r"^## (?P<time>\d+:\d+–\d+:\d+) — .*?^\*\*Narration:\*\*\s*$\s*(?P<body>“.*?”)",
        re.MULTILINE | re.DOTALL,
    )
    return {
        match.group("time"): match.group("body").strip("“”").strip()
        for match in pattern.finditer(text)
    }


def notes_for(narration, *times):
    return "\n\n".join(narration[time] for time in times)


def build_deck():
    narration = extract_narration()
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # Scene 1
    slide = base_slide(
        prs, 1, "0:00–0:35", "Workload-Aware Long-Range Links",
        "Can physical shortcuts follow the workload without violating real constraints?",
    )
    textbox(slide, "16 CHIPLET  •  4×4 MESH", 0.65, 1.83, 4.4, 0.42, 14, MESH, True)
    textbox(
        slide,
        "Distant traffic crosses several intermediate links.",
        0.65, 2.33, 4.4, 1.15, 25, WHITE, True,
    )
    textbox(
        slide,
        "A long-range link can reduce hops—but consumes wire and PHY endpoints.",
        0.65, 3.65, 4.35, 1.05, 16, MUTED,
    )
    draw_mesh(slide, 7.0, 2.05, 3.8, highlight=[0, 4, 8, 12, 13, 14, 15], shortcut=(0, 15))
    pill(slide, "MULTI-HOP ROUTE", 7.05, 6.28, 1.60, CORAL, WHITE)
    pill(slide, "POTENTIAL SHORTCUT", 9.00, 6.28, 1.85, AWARE)
    add_notes(slide, notes_for(narration, "0:00–0:35"))

    # Scene 2
    slide = base_slide(
        prs, 2, "0:35–1:10", "Starting Point vs. Our Implementation",
        "Make the project contribution explicit.",
    )
    pill(slide, "EXISTING TOOLS", 0.64, 1.75, 1.50, MESH, BG)
    existing = [
        ("RapidChiplet", "Physical + analytical\nchiplet modeling", MESH),
        ("BookSim", "Detailed network\nsimulation", FIXED),
        ("STAGE", "Distributed-AI\ntrace generation", PURPLE),
    ]
    for index, (title, body, color) in enumerate(existing):
        x = 0.64 + index * 2.43
        rect(slide, x, 2.20, 2.12, 1.30, PANEL, True, color, 2)
        textbox(slide, title, x + 0.14, 2.35, 1.84, 0.34, 17, color, True, PP_ALIGN.CENTER)
        textbox(slide, body, x + 0.12, 2.75, 1.88, 0.54, 11, WHITE, False, PP_ALIGN.CENTER)
    line(slide, 8.10, 1.70, 8.10, 6.60, GRID, 2)
    pill(slide, "OUR IMPLEMENTATION", 8.48, 1.75, 1.83, AWARE)
    ours = [
        "Reproducible\nworkloads", "8-PHY physical\nbaseline", "96-candidate\nevaluation",
        "Aware + Fixed\nselectors", "STAGE traffic\nconversion", "BookSim + cost\ncomparison",
    ]
    for index, label in enumerate(ours):
        row, col = divmod(index, 2)
        x = 8.48 + col * 2.05
        y = 2.22 + row * 1.35
        rect(slide, x, y, 1.78, 1.03, PANEL_2, True, AWARE if index in (2, 3) else GRID, 1.5)
        textbox(slide, label, x + 0.08, y + 0.12, 1.62, 0.78, 12, WHITE, index in (2, 3), PP_ALIGN.CENTER)
    add_notes(slide, notes_for(narration, "0:35–1:10"))

    # Scene 3
    slide = base_slide(
        prs, 3, "1:10–2:05", "Goal + Reproducible Workloads",
        "OUR IMPLEMENTATION  •  Same traffic matrix for every topology policy.",
    )
    constraints = [("K", "max links"), ("B", "wire budget"), ("PHY", "available endpoints"), ("MaxLoad", "≤ Mesh baseline")]
    for index, (symbol, label) in enumerate(constraints):
        x = 0.64 + index * 3.05
        rect(slide, x, 1.72, 2.72, 0.75, PANEL, True, [AWARE, FIXED, PURPLE, CORAL][index], 1.5)
        textbox(slide, symbol, x + 0.12, 1.82, 0.80, 0.30, 18, [AWARE, FIXED, PURPLE, CORAL][index], True)
        textbox(slide, label, x + 0.92, 1.82, 1.65, 0.30, 12, WHITE, True)
    workloads = [
        ("random_uniform", MESH, [(0, 15), (3, 8), (12, 6)]),
        ("transpose", AWARE, [(1, 4), (3, 12), (6, 9)]),
        ("permutation", PURPLE, [(2, 13), (5, 11), (7, 8)]),
        ("hotspot", CORAL, [(0, 10), (3, 10), (15, 10)]),
    ]
    for index, (name, color, edges) in enumerate(workloads):
        x = 0.64 + index * 3.05
        rect(slide, x, 2.82, 2.72, 3.55, PANEL, True)
        textbox(slide, name, x + 0.12, 3.02, 2.48, 0.35, 14, color, True, PP_ALIGN.CENTER)
        positions = draw_mesh(slide, x + 0.62, 3.82, 1.46, labels=False)
        for u, v in edges:
            line(slide, *positions[u], *positions[v], color, 2.5)
        textbox(
            slide,
            ["broadly distributed", "structured mapping", "one-to-one pairs", "concentrated targets"][index],
            x + 0.14, 5.78, 2.44, 0.35, 10, MUTED, False, PP_ALIGN.CENTER,
        )
    add_notes(slide, notes_for(narration, "1:10–1:30", "1:30–2:05"))

    # Scene 4
    slide = base_slide(
        prs, 4, "2:05–3:05", "From Manual Shortcut to Physical Feasibility",
        "OUR IMPLEMENTATION  •  Exact routing, wire length, and PHY endpoints.",
    )
    rect(slide, 0.62, 1.70, 5.75, 4.95, PANEL, True)
    textbox(slide, "MANUAL FEASIBILITY TEST", 0.86, 1.91, 3.0, 0.35, 13, AWARE, True)
    draw_mesh(slide, 1.45, 2.80, 3.50, shortcut=(3, 12))
    pill(slide, "3 ↔ 12", 2.43, 6.08, 1.25, AWARE)
    textbox(slide, "Add link  →  regenerate SPLIF  →  simulate", 0.92, 6.31, 5.10, 0.24, 10, MUTED, False, PP_ALIGN.CENTER)
    rect(slide, 6.72, 1.70, 5.98, 4.95, PANEL, True)
    textbox(slide, "8-PHY SHORTCUT-CAPABLE CHIPLET", 6.98, 1.91, 4.25, 0.35, 13, FIXED, True)
    cx, cy = 9.67, 4.15
    rect(slide, cx - 0.90, cy - 0.90, 1.80, 1.80, PANEL_2, True, WHITE, 2)
    textbox(slide, "CHIPLET", cx - 0.65, cy - 0.22, 1.30, 0.44, 18, WHITE, True, PP_ALIGN.CENTER)
    phy_positions = [
        (cx - 0.75, cy - 1.12), (cx, cy - 1.12), (cx + 0.75, cy - 1.12),
        (cx + 1.12, cy), (cx + 0.75, cy + 1.12), (cx, cy + 1.12),
        (cx - 0.75, cy + 1.12), (cx - 1.12, cy),
    ]
    for phy, (px, py) in enumerate(phy_positions):
        color = MESH if phy in (1, 3, 5, 7) else AWARE
        rect(slide, px - 0.17, py - 0.17, 0.34, 0.34, color, True)
        textbox(slide, str(phy), px - 0.17, py - 0.17, 0.34, 0.34, 9, BG, True, PP_ALIGN.CENTER, margin=0)
    pill(slide, "CARDINAL: 1 / 3 / 5 / 7", 7.28, 6.08, 2.15, MESH)
    pill(slide, "SPARE: 0 / 2 / 4 / 6", 9.74, 6.08, 2.05, AWARE)
    add_notes(slide, notes_for(narration, "2:05–2:35", "2:35–3:05"))

    # Scene 5
    slide = base_slide(
        prs, 5, "3:05–4:30", "Candidate Evaluation + Workload-Aware Selection",
        "OUR IMPLEMENTATION  •  Reevaluate after every permanent link.",
    )
    pill(slide, "C(16,2) − 24 = 96 CANDIDATES", 0.65, 1.66, 2.53, AWARE)
    stages = [
        ("PAIR", "non-mesh"),
        ("PHY", "shortest free"),
        ("WIRE", "length + cycles"),
        ("ROUTE", "regenerate SPLIF"),
        ("METRICS", "latency + load"),
        ("FILTER", "budget + cap"),
        ("CHOOSE", "min latency"),
    ]
    for index, (head, body) in enumerate(stages):
        x = 0.54 + index * 1.82
        rect(slide, x, 2.30, 1.52, 1.18, PANEL, True, AWARE if index == 6 else GRID, 1.5)
        textbox(slide, head, x + 0.08, 2.47, 1.36, 0.28, 13, AWARE if index == 6 else WHITE, True, PP_ALIGN.CENTER)
        textbox(slide, body, x + 0.08, 2.83, 1.36, 0.33, 9, MUTED, False, PP_ALIGN.CENTER)
        if index < len(stages) - 1:
            textbox(slide, "→", x + 1.53, 2.64, 0.27, 0.32, 18, AWARE, True, PP_ALIGN.CENTER)
    rect(slide, 0.64, 4.05, 7.05, 2.18, PANEL, True)
    textbox(slide, "OBJECTIVE + CONSTRAINTS", 0.90, 4.27, 2.40, 0.30, 12, MESH, True)
    textbox(slide, "min  Lavg(T, Eextra)", 0.92, 4.79, 2.70, 0.45, 24, WHITE, True)
    textbox(slide, "Σ Length(e) ≤ B", 3.78, 4.62, 1.72, 0.36, 15, FIXED, True, PP_ALIGN.CENTER)
    textbox(slide, "|Eextra| ≤ K", 5.62, 4.62, 1.62, 0.36, 15, PURPLE, True, PP_ALIGN.CENTER)
    textbox(slide, "MaxLoad ≤ MaxLoadbaseline", 3.78, 5.23, 3.46, 0.38, 15, CORAL, True, PP_ALIGN.CENTER)
    textbox(slide, "+ valid, unused PHY endpoints", 0.92, 5.50, 2.65, 0.34, 12, AWARE, True)
    rect(slide, 8.06, 4.05, 4.64, 2.18, PANEL_2, True, AWARE, 1.5)
    textbox(slide, "GREEDY K=2 / K=4 LOOP", 8.33, 4.27, 3.98, 0.30, 12, AWARE, True, PP_ALIGN.CENTER)
    loop = ["regenerate routing", "recompute loads", "occupy PHYs", "update wire", "reevaluate"]
    for index, item in enumerate(loop):
        y = 4.76 + index * 0.27
        textbox(slide, f"{index + 1}.  {item}", 8.50, y, 3.50, 0.24, 11, WHITE)
    add_notes(slide, notes_for(narration, "3:05–3:45", "3:45–4:30"))

    # Scene 6
    slide = base_slide(
        prs, 6, "4:30–5:15", "Fair Fixed Baseline + STAGE Integration",
        "OUR IMPLEMENTATION  •  Same resources; only selection traffic changes.",
    )
    rect(slide, 0.62, 1.72, 12.05, 2.03, PANEL, True)
    textbox(slide, "UNIFORM REFERENCE", 0.92, 2.04, 2.10, 0.34, 13, FIXED, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 3.08, 2.03, 0.50, 0.34, 20, FIXED, True, PP_ALIGN.CENTER)
    textbox(slide, "SELECT ONCE", 3.65, 2.04, 1.75, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 5.48, 2.03, 0.50, 0.34, 20, FIXED, True, PP_ALIGN.CENTER)
    textbox(slide, "FREEZE ACROSS WORKLOADS", 6.05, 2.04, 2.80, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    pill(slide, "FIXED", 9.20, 2.02, 1.00, FIXED)
    textbox(slide, "CURRENT WORKLOAD", 0.92, 2.82, 2.10, 0.34, 13, AWARE, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 3.08, 2.81, 0.50, 0.34, 20, AWARE, True, PP_ALIGN.CENTER)
    textbox(slide, "SELECT FOR TRAFFIC", 3.65, 2.82, 2.20, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    pill(slide, "WORKLOAD-AWARE", 6.25, 2.80, 1.82, AWARE)
    textbox(slide, "Same K  •  same B  •  same PHYs  •  same routing  •  same load cap", 8.33, 2.82, 3.85, 0.34, 10, MUTED, True, PP_ALIGN.CENTER)
    rect(slide, 0.62, 4.18, 12.05, 2.15, PANEL_2, True, PURPLE, 1.5)
    pill(slide, "EXISTING", 0.91, 4.48, 0.90, PURPLE)
    textbox(slide, "STAGE TRACE", 2.12, 4.47, 1.65, 0.34, 14, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 3.83, 4.46, 0.50, 0.34, 20, PURPLE, True, PP_ALIGN.CENTER)
    pill(slide, "OUR CONVERSION", 4.42, 4.48, 1.55, AWARE)
    textbox(slide, "validated traffic", 6.18, 4.47, 1.70, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 7.92, 4.46, 0.50, 0.34, 20, AWARE, True, PP_ALIGN.CENTER)
    textbox(slide, "16 CHIPLET MAP", 8.48, 4.47, 1.78, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, "→", 10.31, 4.46, 0.50, 0.34, 20, AWARE, True, PP_ALIGN.CENTER)
    textbox(slide, "SAME SELECTOR", 10.82, 4.47, 1.50, 0.34, 13, AWARE, True, PP_ALIGN.CENTER)
    textbox(slide, "No special optimizer for STAGE", 4.32, 5.35, 4.62, 0.44, 20, WHITE, True, PP_ALIGN.CENTER)
    add_notes(slide, notes_for(narration, "4:30–4:55", "4:55–5:15"))

    # Scene 7
    slide = base_slide(
        prs, 7, "5:15–5:40", "Same Budget. Different Workload. Different Link.",
        "B45  •  requested K=1  •  Workload-Aware selections",
    )
    cards = [
        ("random_uniform", (1, 14), MESH),
        ("transpose", (3, 12), AWARE),
        ("permutation", (4, 14), PURPLE),
        ("hotspot", (4, 11), CORAL),
    ]
    for index, (name, pair, color) in enumerate(cards):
        x = 0.58 + index * 3.16
        rect(slide, x, 1.77, 2.90, 4.60, PANEL, True, color, 2)
        textbox(slide, name, x + 0.12, 1.98, 2.66, 0.34, 14, color, True, PP_ALIGN.CENTER)
        draw_mesh(slide, x + 0.73, 3.00, 1.45, shortcut=pair, labels=False)
        textbox(slide, f"{pair[0]} ↔ {pair[1]}", x + 0.35, 5.31, 2.20, 0.54, 24, color, True, PP_ALIGN.CENTER)
    pill(slide, "FIXED", 4.62, 6.56, 0.95, FIXED)
    textbox(slide, "uniform reference selects 1 ↔ 14", 5.70, 6.54, 3.18, 0.34, 13, WHITE, True, PP_ALIGN.CENTER)
    add_notes(slide, notes_for(narration, "5:15–5:40"))

    # Scene 8
    slide = base_slide(
        prs, 8, "5:40–6:35", "Strong Gain + Diminishing Return",
        "Transpose  •  validated low-load BookSim results",
    )
    rect(slide, 0.54, 1.63, 7.17, 4.88, WHITE, True)
    add_picture_contain(
        slide,
        ROOT / "results/figures/synthetic_aware_gain_vs_fixed_transpose.png",
        0.63, 1.72, 6.99, 4.70,
    )
    rect(slide, 7.96, 1.63, 4.83, 4.88, WHITE, True)
    add_picture_contain(
        slide,
        ROOT / "results/figures/actual_k_vs_requested_k_synthetic.png",
        8.05, 1.72, 4.65, 3.38,
    )
    label_card(slide, "MESH", "98.863 cycles", 0.65, 6.18, 2.05, MESH)
    label_card(slide, "FIXED", "95.076 cycles", 2.86, 6.18, 2.05, FIXED)
    label_card(slide, "WORKLOAD-AWARE", "71.543 cycles", 5.07, 6.18, 2.47, AWARE)
    textbox(slide, "B45 transpose", 8.30, 5.21, 1.75, 0.30, 12, BG, True, PP_ALIGN.CENTER)
    textbox(slide, "requested K", 8.29, 5.62, 1.28, 0.30, 11, GRID, True)
    textbox(slide, "1    2    4", 9.50, 5.62, 1.72, 0.30, 14, BG, True)
    textbox(slide, "actual_k", 8.29, 6.02, 1.28, 0.30, 11, GRID, True)
    textbox(slide, "1    1    1", 9.50, 6.02, 1.72, 0.30, 14, AWARE, True)
    pill(slide, "24.752% vs Fixed", 10.92, 6.04, 1.56, AWARE)
    add_notes(slide, notes_for(narration, "5:40–6:10", "6:10–6:35"))

    # Scene 9
    slide = base_slide(
        prs, 9, "6:35–7:20", "STAGE Result + Final Answer",
        "gpt_pipeline  •  B45  •  requested K=4",
    )
    rect(slide, 0.54, 1.61, 7.20, 4.94, WHITE, True)
    add_picture_contain(
        slide,
        ROOT / "results/figures/stage_latency_comparison_gpt_pipeline.png",
        0.64, 1.72, 7.00, 4.70,
    )
    label_card(slide, "MESH", "45.888 cycles", 7.98, 1.71, 2.20, MESH)
    label_card(slide, "FIXED", "46.085 cycles", 10.38, 1.71, 2.20, FIXED)
    label_card(slide, "WORKLOAD-AWARE", "40.818 cycles", 7.98, 2.82, 4.60, AWARE)
    textbox(slide, "11.428% better than Fixed", 8.28, 3.83, 4.00, 0.44, 19, AWARE, True, PP_ALIGN.CENTER)
    takeaways = [
        ("1", "Topology should follow traffic."),
        ("2", "Physical constraints determine actual links."),
        ("3", "Workload-Aware usually—not always—wins."),
    ]
    for index, (num, text) in enumerate(takeaways):
        y = 4.52 + index * 0.58
        rect(slide, 8.06, y, 0.40, 0.40, [AWARE, FIXED, CORAL][index], True)
        textbox(slide, num, 8.06, y, 0.40, 0.40, 11, BG, True, PP_ALIGN.CENTER, margin=0)
        textbox(slide, text, 8.62, y, 3.83, 0.40, 13, WHITE, True)
    pill(slide, "Fixed wins 5 / 36 STAGE cases", 8.28, 6.42, 3.03, CORAL, WHITE)
    add_notes(slide, notes_for(narration, "6:35–7:00", "7:00–7:20"))

    # Core document metadata.
    prs.core_properties.title = "Workload-Aware Long-Range Link Selection"
    prs.core_properties.subject = "Nine-scene visual deck for the narrated project video"
    prs.core_properties.author = "RapidChiplet project"
    prs.core_properties.keywords = "chiplets, RapidChiplet, BookSim, workload-aware, STAGE"
    prs.core_properties.comments = "Quantitative values are sourced from validated final project results."

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main():
    output = build_deck()
    print("PROJECT VIDEO SLIDES COMPLETE")
    print(f"Slides: 9")
    print(f"Presentation: {output.relative_to(ROOT).as_posix()}")
    print(f"Narration: {SCRIPT_PATH.relative_to(ROOT).as_posix()}")


if __name__ == "__main__":
    main()
